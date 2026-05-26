import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import os
from datetime import date
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import subprocess
import sys

AZUL = "#1a4f8a"
AZUL_CLARO = "#e8f0fb"
BLANCO = "#ffffff"
GRIS = "#f4f6f9"
TEXTO = "#2c3e50"
ACENTO = "#2980b9"

def set_celda_color(celda, hex_color):
    tc = celda._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def agregar_linea_horizontal(doc, color="1a4f8a", grosor=18):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(grosor))
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p

def set_fuente(run, nombre="Calibri", tamaño=11, negrita=False, color_hex=None, italica=False):
    run.font.name = nombre
    run.font.size = Pt(tamaño)
    run.bold = negrita
    run.italic = italica
    if color_hex:
        r, g, b = int(color_hex[0:2],16), int(color_hex[2:4],16), int(color_hex[4:6],16)
        run.font.color.rgb = RGBColor(r, g, b)

# ──────────────────────────────────────────────
#  GENERADORES
# ──────────────────────────────────────────────

def generar_certificado(datos, ruta):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width  = Cm(21.59)
    sec.page_height = Cm(27.94)
    sec.top_margin = sec.bottom_margin = Cm(3)
    sec.left_margin = sec.right_margin = Cm(3)

    # Encabezado empresa
    cab = doc.add_paragraph()
    cab.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cab.add_run(datos.get("empresa", "NOMBRE DE LA EMPRESA").upper())
    set_fuente(r, tamaño=18, negrita=True, color_hex="1a4f8a")

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = sub.add_run(datos.get("nit", "NIT: 000.000.000-0"))
    set_fuente(r2, tamaño=10, color_hex="555555")

    agregar_linea_horizontal(doc, "1a4f8a", 24)

    # Título
    doc.add_paragraph()
    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rt = titulo.add_run("CERTIFICADO LABORAL")
    set_fuente(rt, tamaño=16, negrita=True, color_hex="1a4f8a")

    doc.add_paragraph()

    # Cuerpo
    ciudad_fecha = doc.add_paragraph()
    ciudad_fecha.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rf = ciudad_fecha.add_run(f"{datos.get('ciudad','Bogotá')}, {date.today().strftime('%d de %B de %Y')}")
    set_fuente(rf, tamaño=11, color_hex="555555")

    doc.add_paragraph()

    cuerpo = doc.add_paragraph()
    cuerpo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    texto = (
        f"Quien suscribe, {datos.get('firmante','[Nombre del firmante]')}, en su calidad de "
        f"{datos.get('cargo_firmante','Representante Legal')}, certifica que:\n\n"
        f"El/La señor(a) {datos.get('nombre_empleado','[Nombre del empleado]')}, identificado(a) con "
        f"{datos.get('tipo_doc','C.C.')} N.° {datos.get('num_doc','[Número]')}, labora en nuestra empresa "
        f"desde el {datos.get('fecha_ingreso','[Fecha de ingreso]')}, desempeñando el cargo de "
        f"{datos.get('cargo','[Cargo]')}, con un salario mensual de {datos.get('salario','[Salario]')}."
    )
    run_cuerpo = cuerpo.add_run(texto)
    set_fuente(run_cuerpo, tamaño=11)

    doc.add_paragraph()
    obs = doc.add_paragraph()
    obs.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_obs = obs.add_run(
        f"Se expide la presente certificación a solicitud del interesado para los fines que estime convenientes."
    )
    set_fuente(r_obs, tamaño=11)

    doc.add_paragraph()
    doc.add_paragraph()

    firma = doc.add_paragraph()
    firma.alignment = WD_ALIGN_PARAGRAPH.CENTER
    agregar_linea_horizontal(doc, "888888", 6)
    r_firma = firma.add_run(datos.get("firmante","[Nombre del firmante]"))
    set_fuente(r_firma, negrita=True, tamaño=11)

    cargo_p = doc.add_paragraph()
    cargo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cargo = cargo_p.add_run(datos.get("cargo_firmante","[Cargo]"))
    set_fuente(r_cargo, tamaño=10, color_hex="555555")

    doc.save(ruta)

def generar_factura(datos, ruta):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width  = Cm(21.59)
    sec.page_height = Cm(27.94)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.left_margin = sec.right_margin = Cm(2.5)

    # Encabezado
    tabla_cab = doc.add_table(rows=1, cols=2)
    tabla_cab.style = 'Table Grid'
    c1 = tabla_cab.cell(0,0)
    c2 = tabla_cab.cell(0,1)
    set_celda_color(c1, "1a4f8a")
    set_celda_color(c2, "1a4f8a")

    p1 = c1.paragraphs[0]
    r1 = p1.add_run(datos.get("empresa","MI EMPRESA SAS").upper())
    set_fuente(r1, tamaño=14, negrita=True, color_hex="FFFFFF")
    p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r1b = c1.add_paragraph().add_run(f"NIT: {datos.get('nit','000.000.000-0')}")
    set_fuente(r1b, tamaño=9, color_hex="DDDDDD")

    p2 = c2.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r2 = p2.add_run("FACTURA DE VENTA")
    set_fuente(r2, tamaño=13, negrita=True, color_hex="FFFFFF")
    p_num = c2.add_paragraph()
    p_num.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r2b = p_num.add_run(f"N.° {datos.get('num_factura','001-2025')}")
    set_fuente(r2b, tamaño=10, color_hex="DDDDDD")

    doc.add_paragraph()

    # Datos cliente
    tabla_info = doc.add_table(rows=1, cols=2)
    tabla_info.style = 'Table Grid'
    ci1 = tabla_info.cell(0,0)
    ci2 = tabla_info.cell(0,1)

    def fila_dato(celda, label, valor):
        p = celda.add_paragraph()
        rl = p.add_run(f"{label}: ")
        set_fuente(rl, negrita=True, tamaño=10, color_hex="1a4f8a")
        rv = p.add_run(valor)
        set_fuente(rv, tamaño=10)

    fila_dato(ci1, "Cliente", datos.get("cliente","[Nombre cliente]"))
    fila_dato(ci1, "NIT/CC", datos.get("nit_cliente","[Documento]"))
    fila_dato(ci2, "Fecha", datos.get("fecha", date.today().strftime("%d/%m/%Y")))
    fila_dato(ci2, "Ciudad", datos.get("ciudad","Bogotá"))

    doc.add_paragraph()

    # Tabla de items
    headers = ["#", "Descripción", "Cant.", "Valor Unit.", "Total"]
    anchos = [500, 4200, 800, 1400, 1400]
    tabla_items = doc.add_table(rows=1, cols=5)
    tabla_items.style = 'Table Grid'
    for i, (h, ancho) in enumerate(zip(headers, anchos)):
        c = tabla_items.cell(0, i)
        set_celda_color(c, "1a4f8a")
        p = c.paragraphs[0]
        r = p.add_run(h)
        set_fuente(r, negrita=True, tamaño=10, color_hex="FFFFFF")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    items = datos.get("items", [{"desc":"Producto/Servicio","cant":"1","valor":"0"}])
    total_general = 0
    for idx, item in enumerate(items, 1):
        fila = tabla_items.add_row()
        try:
            total = int(str(item.get("cant","1")).replace(",","")) * int(str(item.get("valor","0")).replace(",","").replace(".",""))
        except:
            total = 0
        total_general += total
        vals = [str(idx), item.get("desc",""), str(item.get("cant","1")), f"${item.get('valor','0')}", f"${total:,}"]
        aligns = [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT,
                  WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT]
        for j, (v, al) in enumerate(zip(vals, aligns)):
            c = fila.cells[j]
            p = c.paragraphs[0]
            r = p.add_run(v)
            set_fuente(r, tamaño=10)
            p.alignment = al

    # Fila total
    fila_total = tabla_items.add_row()
    c_label = fila_total.cells[3]
    set_celda_color(c_label, "e8f0fb")
    pl = c_label.paragraphs[0]
    rl = pl.add_run("TOTAL")
    set_fuente(rl, negrita=True, tamaño=11, color_hex="1a4f8a")
    pl.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    c_total = fila_total.cells[4]
    set_celda_color(c_total, "1a4f8a")
    pt = c_total.paragraphs[0]
    rt = pt.add_run(f"${total_general:,}")
    set_fuente(rt, negrita=True, tamaño=11, color_hex="FFFFFF")
    pt.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_paragraph()
    obs_p = doc.add_paragraph()
    ro = obs_p.add_run(f"Observaciones: {datos.get('observaciones','')}")
    set_fuente(ro, tamaño=10, italica=True, color_hex="555555")

    doc.save(ruta)

def generar_carta(datos, ruta):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width  = Cm(21.59)
    sec.page_height = Cm(27.94)
    sec.top_margin = sec.bottom_margin = Cm(3)
    sec.left_margin = sec.right_margin = Cm(3.5)

    # Membrete
    mem = doc.add_paragraph()
    mem.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = mem.add_run(datos.get("empresa","NOMBRE DE LA EMPRESA").upper())
    set_fuente(r, tamaño=14, negrita=True, color_hex="1a4f8a")
    mem2 = doc.add_paragraph()
    r2 = mem2.add_run(f"{datos.get('direccion','')}  |  {datos.get('telefono','')}  |  {datos.get('email','')}")
    set_fuente(r2, tamaño=9, color_hex="777777")
    agregar_linea_horizontal(doc, "1a4f8a", 18)

    doc.add_paragraph()

    # Fecha y ciudad
    fecha_p = doc.add_paragraph()
    fecha_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rf = fecha_p.add_run(f"{datos.get('ciudad','Bogotá')}, {date.today().strftime('%d de %B de %Y')}")
    set_fuente(rf, tamaño=11)

    doc.add_paragraph()

    # Destinatario
    for linea in [
        datos.get("destinatario","Señor(a):"),
        datos.get("cargo_dest",""),
        datos.get("empresa_dest",""),
        datos.get("ciudad_dest","Ciudad"),
    ]:
        if linea:
            p = doc.add_paragraph()
            r = p.add_run(linea)
            set_fuente(r, tamaño=11)

    doc.add_paragraph()

    # Asunto
    asunto_p = doc.add_paragraph()
    ra = asunto_p.add_run("Asunto: ")
    set_fuente(ra, negrita=True, tamaño=11)
    ra2 = asunto_p.add_run(datos.get("asunto","[Asunto de la carta]"))
    set_fuente(ra2, tamaño=11)

    doc.add_paragraph()

    # Saludo
    sal = doc.add_paragraph()
    rs = sal.add_run(f"Estimado(a) {datos.get('destinatario','señor(a)')},")
    set_fuente(rs, tamaño=11)

    doc.add_paragraph()

    # Cuerpo
    cuerpo_p = doc.add_paragraph()
    cuerpo_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rc = cuerpo_p.add_run(datos.get("cuerpo","[Escriba aquí el cuerpo de la carta.]"))
    set_fuente(rc, tamaño=11)

    doc.add_paragraph()
    doc.add_paragraph()

    # Despedida
    desp = doc.add_paragraph()
    rd = desp.add_run("Atentamente,")
    set_fuente(rd, tamaño=11)

    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph()

    firma = doc.add_paragraph()
    rf2 = firma.add_run(datos.get("firmante","[Nombre del firmante]"))
    set_fuente(rf2, negrita=True, tamaño=11)
    cargo_p = doc.add_paragraph()
    rcp = cargo_p.add_run(datos.get("cargo_firmante","[Cargo]"))
    set_fuente(rcp, tamaño=10, color_hex="555555")

    doc.save(ruta)

def generar_contrato(datos, ruta):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width  = Cm(21.59)
    sec.page_height = Cm(27.94)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.left_margin = sec.right_margin = Cm(3)

    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rt = titulo.add_run(f"CONTRATO DE {datos.get('tipo_contrato','PRESTACIÓN DE SERVICIOS').upper()}")
    set_fuente(rt, tamaño=15, negrita=True, color_hex="1a4f8a")

    agregar_linea_horizontal(doc, "1a4f8a", 18)
    doc.add_paragraph()

    num_p = doc.add_paragraph()
    rn = num_p.add_run(f"Contrato N.°: {datos.get('num_contrato','001-2025')}")
    set_fuente(rn, negrita=True, tamaño=11, color_hex="555555")
    num_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_paragraph()

    def seccion(titulo_sec, contenido):
        t = doc.add_paragraph()
        rt = t.add_run(titulo_sec.upper())
        set_fuente(rt, negrita=True, tamaño=11, color_hex="1a4f8a")
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        rc = p.add_run(contenido)
        set_fuente(rc, tamaño=11)
        doc.add_paragraph()

    seccion("PRIMERA – PARTES",
        f"Entre {datos.get('contratante','[Nombre contratante]')}, identificado con "
        f"{datos.get('doc_contratante','C.C. [Número]')}, en adelante EL CONTRATANTE, y "
        f"{datos.get('contratista','[Nombre contratista]')}, identificado con "
        f"{datos.get('doc_contratista','C.C. [Número]')}, en adelante EL CONTRATISTA.")

    seccion("SEGUNDA – OBJETO",
        datos.get("objeto", "El CONTRATISTA se obliga a prestar los servicios de [descripción del servicio]."))

    seccion("TERCERA – VALOR Y FORMA DE PAGO",
        f"El valor del presente contrato es de {datos.get('valor','[valor en pesos]')}, "
        f"pagaderos de la siguiente manera: {datos.get('forma_pago','[describir forma de pago]')}.")

    seccion("CUARTA – DURACIÓN",
        f"El presente contrato tendrá una duración de {datos.get('duracion','[duración]')}, "
        f"contados a partir del {datos.get('fecha_inicio',date.today().strftime('%d/%m/%Y'))}.")

    seccion("QUINTA – OBLIGACIONES DE LAS PARTES",
        datos.get("obligaciones",
            "EL CONTRATISTA se obliga a: [lista de obligaciones]. "
            "EL CONTRATANTE se obliga a: [lista de obligaciones]."))

    seccion("SEXTA – CLÁUSULA DE CONFIDENCIALIDAD",
        "Las partes se comprometen a guardar estricta confidencialidad sobre la información "
        "que sea de carácter reservado, la cual no podrá ser divulgada sin previo consentimiento escrito.")

    seccion("SÉPTIMA – LEGISLACIÓN APLICABLE",
        f"El presente contrato se regirá por las leyes de la República de Colombia. "
        f"Para todos los efectos legales, las partes se someten a la jurisdicción de los jueces de "
        f"{datos.get('ciudad','Bogotá')}.")

    doc.add_paragraph()
    doc.add_paragraph()

    # Firmas lado a lado
    tabla_firmas = doc.add_table(rows=4, cols=2)
    agregar_linea_horizontal(doc, "888888", 6)

    for i, (nombre, cargo, doc_id) in enumerate([
        (datos.get("contratante","[Contratante]"), "EL CONTRATANTE", datos.get("doc_contratante","")),
        (datos.get("contratista","[Contratista]"), "EL CONTRATISTA", datos.get("doc_contratista","")),
    ]):
        tabla_firmas.cell(0, i).paragraphs[0].add_run("")
        tabla_firmas.cell(1, i).paragraphs[0].add_run("")
        p_nombre = tabla_firmas.cell(2, i).paragraphs[0]
        rn2 = p_nombre.add_run(nombre)
        set_fuente(rn2, negrita=True, tamaño=11)
        p_nombre.alignment = WD_ALIGN_PARAGRAPH.CENTER

        p_cargo = tabla_firmas.cell(3, i).paragraphs[0]
        rc2 = p_cargo.add_run(cargo)
        set_fuente(rc2, tamaño=10, color_hex="555555")
        p_cargo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(ruta)


# ──────────────────────────────────────────────
#  DOCUMENTOS MÉDICOS — HEALTHY SPORT MASSAGE
# ──────────────────────────────────────────────

from docx.oxml.ns import nsmap
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Twips

def _celda_med(celda_obj, color_hex=None, span=None):
    """Aplica color de fondo a celda usando XML directo."""
    if color_hex:
        tc = celda_obj._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), color_hex)
        tcPr.append(shd)
    return celda_obj

def _run_med(parrafo, texto, bold=False, size_pt=9, color=None, italic=False):
    run = parrafo.add_run(texto)
    run.font.name = "Arial"
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color:
        r2, g2, b2 = int(color[0:2],16), int(color[2:4],16), int(color[4:6],16)
        run.font.color.rgb = RGBColor(r2, g2, b2)
    return run

def _p_med(doc_o_celda, texto, bold=False, size_pt=9, align=WD_ALIGN_PARAGRAPH.LEFT,
           color=None, space_before=0, space_after=40, italic=False):
    if hasattr(doc_o_celda, 'add_paragraph'):
        p = doc_o_celda.add_paragraph()
    else:
        p = doc_o_celda
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    _run_med(p, texto, bold=bold, size_pt=size_pt, color=color, italic=italic)
    return p

def _set_col_width(celda_obj, width_cm):
    tc = celda_obj._tc
    tcPr = tc.get_or_add_tcPr()
    tcW = OxmlElement('w:tcW')
    tcW.set(qn('w:w'), str(int(width_cm * 567)))  # 1cm ≈ 567 twips
    tcW.set(qn('w:type'), 'dxa')
    tcPr.append(tcW)

def generar_consentimiento(datos, ruta):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width  = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    sec.left_margin = sec.right_margin = Cm(1.8)

    def tabla(cols, widths, estilo='Table Grid'):
        t = doc.add_table(rows=0, cols=cols)
        t.style = estilo
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        return t

    def fila(t, contenidos, heights=None, colores=None):
        row = t.add_row()
        for i, (c_obj, contenido) in enumerate(zip(row.cells, contenidos)):
            if colores and colores[i]:
                _celda_med(c_obj, colores[i])
            if callable(contenido):
                contenido(c_obj)
            elif isinstance(contenido, str):
                p = c_obj.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                _run_med(p, contenido, size_pt=9)
        return row

    def espacio(n=1):
        for _ in range(n):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)

    # ── Título ──
    espacio()
    t1 = _p_med(doc, "CONSENTIMIENTO INFORMADO", bold=True, size_pt=13,
                align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    _p_med(doc, "TERAPIA FÍSICA", bold=True, size_pt=13,
           align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)

    # ── Tabla FECHA / SEDE / SERVICIO ──
    t = tabla(3, [])
    r = t.add_row()
    for c, txt in zip(r.cells, ["FECHA", "SEDE", "SERVICIO"]):
        p = c.paragraphs[0]
        _run_med(p, txt, bold=True, size_pt=9)
    r2 = t.add_row()
    for i, (c, key) in enumerate(zip(r2.cells, ["fecha", "sede", "servicio"])):
        c.paragraphs[0].add_run(datos.get(key, ""))

    espacio()

    # ── 1. Datos de identificación ──
    _p_med(doc, "1.   DATOS DE IDENTIFICACIÓN", bold=True, size_pt=10, space_after=2)
    t2 = tabla(4, [])
    r = t2.add_row()
    for c, txt in zip(r.cells, ["NOMBRE DEL PACIENTE", "EDAD", "Nº HC", "EPS"]):
        _run_med(c.paragraphs[0], txt, bold=True, size_pt=9)
    r2 = t2.add_row()
    vals = [datos.get("nombre_paciente",""), datos.get("edad",""),
            datos.get("num_hc",""), datos.get("eps","")]
    for c, v in zip(r2.cells, vals):
        _run_med(c.paragraphs[0], v, size_pt=9)

    espacio()

    # ── 2. Información sobre el procedimiento ──
    _p_med(doc, "2.   INFORMACIÓN SOBRE EL PROCEDIMIENTO", bold=True, size_pt=10, space_after=2)

    t3 = tabla(1, [])
    r = t3.add_row()
    _run_med(r.cells[0].paragraphs[0], "A.  NOMBRE TECNICO DEL PROCEDIMIENTO:", bold=True, size_pt=9)
    r2 = t3.add_row()
    _run_med(r2.cells[0].paragraphs[0], "TERAPIA FISICA INTEGRAL", size_pt=9)

    espacio()
    _p_med(doc, "B.  DESCRIPCIÓN.", bold=True, size_pt=9, space_after=2)

    textos_sec2 = [
        ("TERAPIA FÍSICA INTEGRAL: ", True,
         "Conjunto de modalidades de rehabilitación corporal implementadas después de una lesión o enfermedad que se aplican de modo simultaneo con el objetivo de recuperar el movimiento completo a través de técnicas integrales."),
        ("C.  OBJETIVO DE REALIZAR EL PROCEDIMIENTO: ", True,
         "Contribuir en la recuperación, prevención y tratamiento de diferentes trastornos corporales bajo una serie de técnicas y movimientos, con la finalidad de reducir y mejorar los niveles de dolor y así mismo generar una adecuada calidad de vida."),
        ("D.  RIESGOS Y/O COMPLICACIONES: ", True, " Fatiga muscular, inflamación, hematoma por ventosa."),
        ("E.  OTRAS ALTERNATIVAS DISPONIBLES: ", True, "Modalidades manuales y mecánicas."),
        ("F.  RIESGOS DE NO TRATARSE: ", True,
         "Parte del trabajo del fisioterapeuta; es explicarle cómo proteger sus articulaciones, cómo entrenar zonas sensibles a lesiones y cuando ya hay lesiones presentes, qué ejercicios debe potenciar, cuáles evitar y cuándo debe descansar. Por esto la importancia de generar una adaptación física mediante la rehabilitación física."),
    ]
    for label, bold_l, resto in textos_sec2:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        _run_med(p, label, bold=True, size_pt=9)
        _run_med(p, resto, size_pt=9)

    espacio()
    _p_med(doc, "G.  INFORMACIÓN ADICIONAL:", bold=True, size_pt=9, space_after=0)
    _p_med(doc, "     En caso de requerir más información sobre el procedimiento a realizar, diríjase al fisioterapeuta tratante.", size_pt=9, space_after=4)

    espacio()

    # ── 3. Declaración del paciente ──
    _p_med(doc, "3.   DECLARACIÓN DEL PACIENTE", bold=True, size_pt=10, space_after=3)

    for txt in [
        "Me han explicado y he comprendido satisfactoriamente la esencia y el propósito de este procedimiento, también me han aclarado todas las dudas y me han dicho los posibles riesgos y complicaciones, así como las otras alternativas de tratamiento.",
        "Doy mi consentimiento para que me realicen el procedimiento descrito anteriormente y los procedimientos complementarios que sean necesarios o convenientes mediante la realización de este, a criterio de los profesionales que lo llevan a cabo.",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        _run_med(p, txt, size_pt=9)

    espacio()

    t_f1 = tabla(2, [])
    r = t_f1.add_row()
    c_izq, c_der = r.cells[0], r.cells[1]
    p1 = c_izq.paragraphs[0]
    _run_med(p1, f"Firma del paciente: {'_'*30}", size_pt=9)
    c_izq.add_paragraph().add_run("")
    p2 = c_izq.add_paragraph()
    _run_med(p2, f"CC: {'_'*30}", size_pt=9)
    p_der = c_der.paragraphs[0]
    p_der.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run_med(p_der, "(foto)", size_pt=9, color="AAAAAA")

    # ── PÁGINA 2 ──
    doc.add_page_break()

    _p_med(doc, "CONSENTIMIENTO INFORMADO", bold=True, size_pt=13,
           align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    _p_med(doc, "TERAPIA FÍSICA", bold=True, size_pt=13,
           align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)

    _p_med(doc, "4.   DECLARACIÓN DEL RESPONSABLE DEL PACIENTE (Sólo en caso de menor de edad, paciente interdicto y/o analfabetismo, o adulto mayor)",
           bold=True, size_pt=9, space_after=4)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(4)
    _run_med(p,
        f"Yo {'_'*20} sé que el paciente {'_'*20} con Nº de identificación {'_'*20} ha sido considerado "
        "por ahora incapaz de tomar por sí mismo la decisión de aceptar o rechazar el procedimiento. "
        "También se me han explicado los riesgos y complicaciones, así como las otras alternativas de tratamiento. "
        "Soy consciente que no existen garantías absolutas de los resultados del procedimiento. He comprendido todo "
        "lo anterior perfectamente y por ello doy mi consentimiento para que los profesionales tratantes y el "
        "personal auxiliar que precise le realicen este procedimiento", size_pt=9)

    espacio()
    t_f2 = tabla(2, [])
    r2 = t_f2.add_row()
    ci, cd = r2.cells[0], r2.cells[1]
    _run_med(ci.paragraphs[0], f"Firma de quien autoriza: {'_'*30}", size_pt=9)
    ci.add_paragraph().add_run("")
    _run_med(ci.add_paragraph(), f"CC: {'_'*30}", size_pt=9)
    cd.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run_med(cd.paragraphs[0], "(foto)", size_pt=9, color="AAAAAA")

    espacio()
    _p_med(doc, "5.   DECLARACION DEL PROFESIONAL TRATANTE", bold=True, size_pt=10, space_after=3)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(4)
    _run_med(p,
        f"Yo {'_'*25} como profesional tratante, he informado al paciente sobre la esencia y el propósito del "
        "procedimiento descrito anteriormente, de sus alternativas, posibles riesgos, resultados esperados y que "
        "no existen garantías absolutas de los resultados del procedimiento.", size_pt=9)

    espacio()
    t_f3 = tabla(1, [])
    r3 = t_f3.add_row()
    _run_med(r3.cells[0].paragraphs[0], f"Firma del profesional: {'_'*35}", size_pt=9)
    r3.cells[0].add_paragraph().add_run("")
    _run_med(r3.cells[0].add_paragraph(), f"CC: {'_'*35}", size_pt=9)
    r3.cells[0].add_paragraph().add_run("")
    _run_med(r3.cells[0].add_paragraph(),
             f"Cargo: Fisioterapeuta  {'_'*35}", size_pt=9)

    # Texto legal al final
    for _ in range(6):
        espacio()
    p_legal = doc.add_paragraph()
    p_legal.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _run_med(p_legal,
        "Con la firma de este formato, autorizo de manera previa e informada, expresa e inequívoca, para que se dé "
        "tratamiento a mis datos personales, de acuerdo a las finalidades descritas en el manual de políticas de "
        "HEALTHY SPORT MASSAGE., quien actúa como responsable del tratamiento de datos.",
        size_pt=8, color="555555")

    doc.save(ruta)


def generar_historia_clinica(datos, ruta):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width  = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    sec.left_margin = sec.right_margin = Cm(1.5)

    def tabla_grid(cols):
        t = doc.add_table(rows=0, cols=cols)
        t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        return t

    def celda_label(c, txt, shade="DDDDDD"):
        _celda_med(c, shade)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        _run_med(p, txt, bold=True, size_pt=9)

    def celda_valor(c, txt=""):
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        _run_med(p, txt, size_pt=9)

    def espacio(n=1):
        for _ in range(n):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1)

    # ── Encabezado con logo ──
    t_enc = tabla_grid(2)
    r = t_enc.add_row()
    _celda_med(r.cells[0])
    _run_med(r.cells[0].paragraphs[0], "HISTORIA CLINICA HEALTHY SPORT MASSAGE", bold=True, size_pt=13)
    _celda_med(r.cells[1], "F0F0F0")
    p_logo = r.cells[1].paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run_med(p_logo, "Healthy\nSport Massage", bold=True, size_pt=9)

    espacio()

    # ── Datos personales ──
    claves_vals = [
        ("nombre_completo", "documento", "edad"),
        ("eps", "num_contacto", None),
        ("direccion", "correo", None),
        ("barrio", "talla", None),
        ("fecha_nac", "peso", None),
    ]

    t_p2 = tabla_grid(4)
    filas_p = [
        ("NOMBRE COMPLETO",        "nombre_completo",  None,                 None),
        ("DOCUMENTO DE IDENTIDAD", "documento",        "EDAD",               "edad"),
        ("EPS",                    "eps",              "NUMERO DE CONTACTO", "num_contacto"),
        ("DIRECCION DE VIVIENDA",  "direccion",        "CORREO",             "correo"),
        ("BARRIO",                 "barrio",           "TALLA",              "talla"),
        ("FECHA DE NACIMIENTO",    "fecha_nac",        "PESO",               "peso"),
    ]
    for lbl1, key1, lbl2, key2 in filas_p:
        r = t_p2.add_row()
        celda_label(r.cells[0], lbl1)
        celda_valor(r.cells[1], datos.get(key1, "") if key1 else "")
        celda_label(r.cells[2], lbl2 or "")
        celda_valor(r.cells[3], datos.get(key2, "") if key2 else "")

    espacio()

    # ── Antecedentes (4 filas: FAMILIARES, PATOLOGICOS, FARMACOLOGICOS, GINECOLOGICOS) ──
    # Columnas: [ANTECEDENTES label | sublabel | espacio para escribir | espacio extra]
    t_ant = tabla_grid(4)
    sub_ants = ["FAMILIARES", "PATOLOGICOS", "FARMACOLOGICOS", "GINECOLOGICOS"]
    claves_ant = ["ant_familiares", "ant_patologicos", "ant_farmacologicos", "ant_ginecologicos"]
    for i, (sl, key) in enumerate(zip(sub_ants, claves_ant)):
        r = t_ant.add_row()
        # Columna 0: label principal solo en la primera fila, resto vacío gris
        celda_label(r.cells[0], "ANTECEDENTES" if i == 0 else "", shade="CCCCCC")
        # Columna 1: sublabel
        celda_label(r.cells[1], sl, shade="EEEEEE")
        # Columnas 2 y 3: espacio para escribir (mergeadas)
        r.cells[2].merge(r.cells[3])
        celda_valor(r.cells[2], datos.get(key, ""))

    espacio()

    # ── Cirugías (3 filas de espacio para escribir) ──
    t_cir = tabla_grid(2)
    for i in range(3):
        r = t_cir.add_row()
        # Altura mínima para que el espacio sea visible
        celda_label(r.cells[0], "CIRUGIAS" if i == 0 else "", shade="CCCCCC")
        # Agregar altura a la celda de valor para que sea escribible
        c_val = r.cells[1]
        celda_valor(c_val, datos.get("cirugias", "") if i == 0 else "")
        # Agregar párrafos vacíos para dar altura visible
        for _ in range(2):
            p_extra = c_val.add_paragraph()
            p_extra.paragraph_format.space_before = Pt(0)
            p_extra.paragraph_format.space_after = Pt(0)

    espacio()

    # ── Alergias (2 filas de espacio para escribir) ──
    t_al = tabla_grid(2)
    for i in range(2):
        r = t_al.add_row()
        celda_label(r.cells[0], "ALERGIAS" if i == 0 else "", shade="CCCCCC")
        c_val = r.cells[1]
        celda_valor(c_val, datos.get("alergias", "") if i == 0 else "")
        for _ in range(2):
            p_extra = c_val.add_paragraph()
            p_extra.paragraph_format.space_before = Pt(0)
            p_extra.paragraph_format.space_after = Pt(0)

    espacio()

    # ── Motivo de la consulta ──
    t_mot = tabla_grid(2)
    r = t_mot.add_row()
    _celda_med(r.cells[0], "CCCCCC")
    _run_med(r.cells[0].paragraphs[0], "MOTIVO DE LA CONSULTA", bold=True, size_pt=11)
    celda_valor(r.cells[1], datos.get("motivo_consulta", ""))
    for _ in range(9):
        r = t_mot.add_row()
        celda_label(r.cells[0], "", shade="CCCCCC")
        celda_valor(r.cells[1])

    espacio()

    # ── Tratamiento ──
    t_trat = tabla_grid(2)
    r = t_trat.add_row()
    _celda_med(r.cells[0], "CCCCCC")
    _run_med(r.cells[0].paragraphs[0], "TRATAMIENTO", bold=True, size_pt=11)
    celda_valor(r.cells[1], datos.get("tratamiento", ""))
    for _ in range(9):
        r = t_trat.add_row()
        celda_label(r.cells[0], "", shade="CCCCCC")
        celda_valor(r.cells[1])

    # ── PÁGINA 2: Seguimiento ──
    doc.add_page_break()

    t_enc2 = tabla_grid(2)
    r2 = t_enc2.add_row()
    _run_med(r2.cells[0].paragraphs[0], "HISTORIA CLINICA HEALTHY SPORT MASSAGE", bold=True, size_pt=13)
    _celda_med(r2.cells[1], "F0F0F0")
    r2.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run_med(r2.cells[1].paragraphs[0], "Healthy Sport Massage", bold=True, size_pt=9)

    espacio()

    t_seg = tabla_grid(4)
    r_h = t_seg.add_row()
    for c, txt, shade in zip(r_h.cells,
                              ["TERAPIA", "FECHA", "EVOLUCIÓN", "OTROS"],
                              ["CCCCCC","CCCCCC","CCCCCC","CCCCCC"]):
        _celda_med(c, shade)
        p = c.paragraphs[0]
        _run_med(p, txt, bold=True, size_pt=9)

    # Primera fila con Plasma/Suero/Infiltración
    r1 = t_seg.add_row()
    for c in r1.cells[:3]:
        celda_valor(c)
    otros_c = r1.cells[3]
    for sub in ["PLASMA", "SUERO", "INFILTRACIÓN"]:
        p = otros_c.add_paragraph() if sub != "PLASMA" else otros_c.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        _run_med(p, sub, bold=True, size_pt=9)

    for _ in range(27):
        r = t_seg.add_row()
        for c in r.cells:
            celda_valor(c)

    doc.save(ruta)


# ──────────────────────────────────────────────
#  INTERFAZ GRÁFICA
# ──────────────────────────────────────────────

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Generador de Documentos Word")
        self.geometry("820x700")
        self.configure(bg=GRIS)
        self.resizable(True, True)
        self._construir_ui()

    def _construir_ui(self):
        # Header
        header = tk.Frame(self, bg=AZUL, height=70)
        header.pack(fill="x")
        tk.Label(header, text="📄  Generador de Documentos", font=("Segoe UI", 17, "bold"),
                 bg=AZUL, fg=BLANCO).pack(side="left", padx=20, pady=15)

        # Selector de tipo
        selector_frame = tk.Frame(self, bg=GRIS, pady=10)
        selector_frame.pack(fill="x", padx=20)
        tk.Label(selector_frame, text="Tipo de documento:", font=("Segoe UI", 11, "bold"),
                 bg=GRIS, fg=TEXTO).pack(side="left")

        self.tipo_var = tk.StringVar(value="Certificado")
        tipos = ["Certificado", "Factura", "Carta / Oficio", "Contrato",
                 "Consentimiento Informado (HSM)", "Historia Clínica (HSM)"]
        combo = ttk.Combobox(selector_frame, textvariable=self.tipo_var, values=tipos,
                             state="readonly", font=("Segoe UI", 11), width=22)
        combo.pack(side="left", padx=10)
        combo.bind("<<ComboboxSelected>>", lambda e: self._cargar_formulario())

        btn_gen = tk.Button(selector_frame, text="⬇  Generar .docx",
                            font=("Segoe UI", 11, "bold"), bg=AZUL, fg=BLANCO,
                            relief="flat", padx=18, pady=6, cursor="hand2",
                            command=self._generar)
        btn_gen.pack(side="right")

        # Contenedor del formulario con scroll
        container = tk.Frame(self, bg=GRIS)
        container.pack(fill="both", expand=True, padx=20, pady=10)

        canvas = tk.Canvas(container, bg=GRIS, highlightthickness=0)
        scrollbar = ttk.Scrollbar(container, orient="vertical", command=canvas.yview)
        self.scroll_frame = tk.Frame(canvas, bg=GRIS)

        self.scroll_frame.bind("<Configure>",
            lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.create_window((0, 0), window=self.scroll_frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)

        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")
        canvas.bind_all("<MouseWheel>", lambda e: canvas.yview_scroll(-1*(e.delta//120), "units"))

        self.campos = {}
        self._cargar_formulario()

    def _limpiar_formulario(self):
        for w in self.scroll_frame.winfo_children():
            w.destroy()
        self.campos = {}

    def _agregar_campo(self, label, key, alto=1, placeholder=""):
        frame = tk.Frame(self.scroll_frame, bg=GRIS)
        frame.pack(fill="x", pady=4)
        tk.Label(frame, text=label, font=("Segoe UI", 10, "bold"),
                 bg=GRIS, fg=TEXTO, width=22, anchor="w").pack(side="left")
        if alto == 1:
            var = tk.StringVar(value=placeholder)
            entry = tk.Entry(frame, textvariable=var, font=("Segoe UI", 10),
                             bg=BLANCO, relief="solid", bd=1)
            entry.pack(side="left", fill="x", expand=True, ipady=5)
            self.campos[key] = var
        else:
            text = tk.Text(frame, font=("Segoe UI", 10), bg=BLANCO,
                           relief="solid", bd=1, height=alto, wrap="word")
            if placeholder:
                text.insert("1.0", placeholder)
            text.pack(side="left", fill="x", expand=True)
            self.campos[key] = text

    def _separador(self, titulo):
        f = tk.Frame(self.scroll_frame, bg=AZUL_CLARO)
        f.pack(fill="x", pady=(14, 4))
        tk.Label(f, text=f"  {titulo}", font=("Segoe UI", 10, "bold"),
                 bg=AZUL_CLARO, fg=AZUL, pady=4).pack(side="left")

    def _cargar_formulario(self):
        self._limpiar_formulario()
        tipo = self.tipo_var.get()

        if tipo == "Certificado":
            self._separador("Datos de la empresa")
            self._agregar_campo("Empresa", "empresa", placeholder="Mi Empresa SAS")
            self._agregar_campo("NIT", "nit", placeholder="900.123.456-7")
            self._separador("Datos del empleado")
            self._agregar_campo("Nombre empleado", "nombre_empleado", placeholder="Juan Pérez")
            self._agregar_campo("Tipo documento", "tipo_doc", placeholder="C.C.")
            self._agregar_campo("Número documento", "num_doc", placeholder="1.012.345.678")
            self._agregar_campo("Cargo", "cargo", placeholder="Analista de Sistemas")
            self._agregar_campo("Fecha ingreso", "fecha_ingreso", placeholder="01 de enero de 2022")
            self._agregar_campo("Salario mensual", "salario", placeholder="$3.000.000")
            self._separador("Firmante")
            self._agregar_campo("Nombre firmante", "firmante", placeholder="Gerente General")
            self._agregar_campo("Cargo firmante", "cargo_firmante", placeholder="Gerente General")
            self._agregar_campo("Ciudad", "ciudad", placeholder="Bogotá")

        elif tipo == "Factura":
            self._separador("Datos de la empresa")
            self._agregar_campo("Empresa", "empresa", placeholder="Mi Empresa SAS")
            self._agregar_campo("NIT empresa", "nit", placeholder="900.123.456-7")
            self._separador("Datos del cliente")
            self._agregar_campo("Cliente", "cliente", placeholder="Cliente XYZ")
            self._agregar_campo("NIT / CC cliente", "nit_cliente", placeholder="1.000.000.000")
            self._agregar_campo("Ciudad", "ciudad", placeholder="Bogotá")
            self._agregar_campo("Fecha", "fecha", placeholder=date.today().strftime("%d/%m/%Y"))
            self._agregar_campo("N.° Factura", "num_factura", placeholder="001-2025")
            self._separador("Ítems (uno por línea: descripción|cantidad|valor_unitario)")
            self._agregar_campo("Ítems", "items_raw", alto=5,
                placeholder="Servicio de diseño web|1|2500000\nHoras de soporte|5|80000")
            self._agregar_campo("Observaciones", "observaciones",
                placeholder="Gracias por su compra.")

        elif tipo == "Carta / Oficio":
            self._separador("Membrete")
            self._agregar_campo("Empresa / Remitente", "empresa", placeholder="Mi Empresa SAS")
            self._agregar_campo("Dirección", "direccion", placeholder="Calle 10 # 5-20, Bogotá")
            self._agregar_campo("Teléfono", "telefono", placeholder="601 234 5678")
            self._agregar_campo("Email", "email", placeholder="info@miempresa.com")
            self._separador("Destinatario")
            self._agregar_campo("Nombre destinatario", "destinatario", placeholder="Dr. Carlos Ruiz")
            self._agregar_campo("Cargo destinatario", "cargo_dest", placeholder="Director de Compras")
            self._agregar_campo("Empresa destinatario", "empresa_dest", placeholder="Empresa ABC")
            self._agregar_campo("Ciudad destinatario", "ciudad_dest", placeholder="Medellín")
            self._separador("Contenido")
            self._agregar_campo("Ciudad remitente", "ciudad", placeholder="Bogotá")
            self._agregar_campo("Asunto", "asunto", placeholder="Propuesta comercial")
            self._agregar_campo("Cuerpo de la carta", "cuerpo", alto=6,
                placeholder="Por medio de la presente, nos permitimos...")
            self._separador("Firmante")
            self._agregar_campo("Nombre firmante", "firmante", placeholder="Ana López")
            self._agregar_campo("Cargo firmante", "cargo_firmante", placeholder="Gerente Comercial")

        elif tipo == "Contrato":
            self._separador("Identificación del contrato")
            self._agregar_campo("N.° Contrato", "num_contrato", placeholder="001-2025")
            self._agregar_campo("Tipo de contrato", "tipo_contrato", placeholder="Prestación de Servicios")
            self._agregar_campo("Ciudad", "ciudad", placeholder="Bogotá")
            self._separador("Contratante")
            self._agregar_campo("Nombre contratante", "contratante", placeholder="Empresa XYZ SAS")
            self._agregar_campo("Doc. contratante", "doc_contratante", placeholder="NIT 900.000.000-1")
            self._separador("Contratista")
            self._agregar_campo("Nombre contratista", "contratista", placeholder="Juan García")
            self._agregar_campo("Doc. contratista", "doc_contratista", placeholder="C.C. 1.012.345.678")
            self._separador("Cláusulas")
            self._agregar_campo("Objeto del contrato", "objeto", alto=3,
                placeholder="El contratista se obliga a desarrollar...")
            self._agregar_campo("Valor del contrato", "valor", placeholder="$5.000.000")
            self._agregar_campo("Forma de pago", "forma_pago", placeholder="50% al inicio, 50% al finalizar")
            self._agregar_campo("Duración", "duracion", placeholder="3 meses")
            self._agregar_campo("Fecha de inicio", "fecha_inicio",
                placeholder=date.today().strftime("%d/%m/%Y"))
            self._agregar_campo("Obligaciones", "obligaciones", alto=4,
                placeholder="El contratista se obliga a entregar el trabajo en los tiempos pactados...")

        elif tipo == "Consentimiento Informado (HSM)":
            self._separador("Datos del paciente")
            self._agregar_campo("Nombre del paciente", "nombre_paciente", placeholder="Juan Pérez")
            self._agregar_campo("Edad", "edad", placeholder="35")
            self._agregar_campo("N.° Historia Clínica", "num_hc", placeholder="HC-001")
            self._agregar_campo("EPS", "eps", placeholder="Sura / Compensar / etc.")
            self._separador("Encabezado del documento")
            self._agregar_campo("Fecha", "fecha", placeholder=date.today().strftime("%d/%m/%Y"))
            self._agregar_campo("Sede", "sede", placeholder="Sede Centro")
            self._agregar_campo("Servicio", "servicio", placeholder="Terapia Física")

        elif tipo == "Historia Clínica (HSM)":
            self._separador("Datos personales del paciente")
            self._agregar_campo("Nombre completo", "nombre_completo", placeholder="Juan Pérez")
            self._agregar_campo("Documento de identidad", "documento", placeholder="C.C. 1.012.345.678")
            self._agregar_campo("Edad", "edad", placeholder="35")
            self._agregar_campo("EPS", "eps", placeholder="Sura")
            self._agregar_campo("Número de contacto", "num_contacto", placeholder="310 000 0000")
            self._agregar_campo("Dirección de vivienda", "direccion", placeholder="Calle 10 # 5-20")
            self._agregar_campo("Correo electrónico", "correo", placeholder="correo@email.com")
            self._agregar_campo("Barrio", "barrio", placeholder="La Candelaria")
            self._agregar_campo("Talla", "talla", placeholder="1.75 m")
            self._agregar_campo("Fecha de nacimiento", "fecha_nac", placeholder="01/01/1990")
            self._agregar_campo("Peso", "peso", placeholder="70 kg")
            self._separador("Antecedentes")
            self._agregar_campo("Familiares", "ant_familiares", placeholder="Hipertensión, diabetes...")
            self._agregar_campo("Patológicos", "ant_patologicos", placeholder="Artritis, hernias...")
            self._agregar_campo("Farmacológicos", "ant_farmacologicos", placeholder="Ibuprofeno, omeprazol...")
            self._agregar_campo("Ginecológicos", "ant_ginecologicos", placeholder="N/A")
            self._agregar_campo("Cirugías", "cirugias", placeholder="Apendicectomía 2018...")
            self._agregar_campo("Alergias", "alergias", placeholder="Penicilina, látex...")
            self._separador("Consulta (opcional — también puedes llenarlo en Word)")
            self._agregar_campo("Motivo de consulta", "motivo_consulta", alto=3,
                placeholder="Dolor lumbar crónico...")
            self._agregar_campo("Tratamiento", "tratamiento", alto=3,
                placeholder="Masaje terapéutico, electroterapia...")
        datos = {}
        for key, widget in self.campos.items():
            if isinstance(widget, tk.StringVar):
                datos[key] = widget.get().strip()
            elif isinstance(widget, tk.Text):
                datos[key] = widget.get("1.0", "end-1c").strip()
        return datos

    def _generar(self):
        tipo = self.tipo_var.get()
        datos = self._leer_campos()

        nombres = {
            "Certificado": "certificado",
            "Factura": "factura",
            "Carta / Oficio": "carta_oficio",
            "Contrato": "contrato",
            "Consentimiento Informado (HSM)": "consentimiento_informado",
            "Historia Clínica (HSM)": "historia_clinica",
        }
        nombre_sugerido = f"{nombres.get(tipo,'documento')}_{date.today().strftime('%Y%m%d')}.docx"

        ruta = filedialog.asksaveasfilename(
            defaultextension=".docx",
            filetypes=[("Word Document", "*.docx")],
            initialfile=nombre_sugerido,
            title="Guardar documento como..."
        )
        if not ruta:
            return

        try:
            if tipo == "Certificado":
                generar_certificado(datos, ruta)
            elif tipo == "Factura":
                # Parsear items
                items = []
                for linea in datos.get("items_raw","").split("\n"):
                    partes = linea.split("|")
                    if len(partes) >= 3:
                        items.append({"desc": partes[0], "cant": partes[1], "valor": partes[2]})
                    elif len(partes) == 1 and partes[0]:
                        items.append({"desc": partes[0], "cant": "1", "valor": "0"})
                datos["items"] = items if items else [{"desc":"Ítem", "cant":"1", "valor":"0"}]
                generar_factura(datos, ruta)
            elif tipo == "Carta / Oficio":
                generar_carta(datos, ruta)
            elif tipo == "Contrato":
                generar_contrato(datos, ruta)
            elif tipo == "Consentimiento Informado (HSM)":
                generar_consentimiento(datos, ruta)
            elif tipo == "Historia Clínica (HSM)":
                generar_historia_clinica(datos, ruta)

            messagebox.showinfo("✅ Documento generado",
                f"Documento guardado exitosamente en:\n{ruta}")
            # Abrir el archivo automáticamente
            if sys.platform == "win32":
                os.startfile(ruta)
            elif sys.platform == "darwin":
                subprocess.Popen(["open", ruta])
            else:
                subprocess.Popen(["xdg-open", ruta])

        except Exception as e:
            messagebox.showerror("Error", f"No se pudo generar el documento:\n{str(e)}")


if __name__ == "__main__":
    app = App()
    app.mainloop()
